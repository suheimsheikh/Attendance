import re, openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

def norm(s):
    return re.sub(r'[^a-z0-9]', '', str(s or '').lower())

# ---------- 1) Parse the weekly diet (calc.xlsx) ----------
wb = openpyxl.load_workbook('/app/menu_src/calc.xlsx', data_only=True)
ws = wb['Sheet1']
rows = list(ws.iter_rows(values_only=True))

# day blocks: (label, start_col, has_meal, has_fiber)
blocks = [
    ("Monday", 0, True, True), ("Tuesday", 8, True, True),
    ("Wednesday", 16, True, True), ("Thursday", 24, True, True),
    ("Friday", 32, True, True), ("Saturday", 40, True, True),
    ("Sunday", 48, True, True), ("Extras (elite)", 56, False, False),
]
diet = []  # dict rows
for label, c, has_meal, has_fib in blocks:
    meal = ""
    for r in rows[3:]:
        if has_meal:
            m = r[c]; menu = r[c+1]
            e,ca,p,f,fib = r[c+2], r[c+3], r[c+4], r[c+5], r[c+6]
        else:
            m = None; menu = r[c]
            e,ca,p,f = r[c+1], r[c+2], r[c+3], r[c+4]; fib = None
        if m and str(m).strip():
            meal = str(m).strip()
        if not menu or not str(menu).strip():
            continue
        if str(menu).strip().lower() == "total":
            continue
        diet.append({"day": label, "meal": meal if has_meal else "",
                     "dish": str(menu).strip(),
                     "energy": e, "carbs": ca, "protein": p, "fat": f, "fiber": fib})

# ---------- 2) Parse recipes (recipes.xlsx) ----------
wb2 = openpyxl.load_workbook('/app/menu_src/recipes.xlsx', data_only=True)
ws2 = wb2['Sheet3']
rrows = list(ws2.iter_rows(values_only=True))
recipes = {}      # norm_name -> {"name":..., "ingredients":[(ing,qty)]}
cur = None
for r in rrows[2:]:
    meal, sl, menu, ing, qty = (r+ (None,)*5)[:5]
    if menu and str(menu).strip():
        cur = str(menu).strip()
        # register full name + each '&'-split part
        for part in [cur] + re.split(r'\s*&\s*', cur):
            k = norm(part)
            if k:
                recipes.setdefault(k, {"name": cur, "ingredients": []})
    if cur and ing and str(ing).strip():
        entry = recipes.get(norm(cur))
        if entry is not None:
            entry["ingredients"].append((str(ing).strip(), str(qty or "").strip()))

# ---------- 3) Match & build combined rows ----------
combined = []
diet_no_recipe = []
matched_recipe_keys = set()
for d in diet:
    key = norm(d["dish"])
    rec = recipes.get(key)
    if rec is None:
        # try each '&'-split part of the diet dish
        for part in re.split(r'\s*&\s*', d["dish"]):
            if norm(part) in recipes:
                rec = recipes[norm(part)]; key = norm(part); break
    if rec is None:
        diet_no_recipe.append(d["dish"])
        combined.append({**d, "ingredient": "— no recipe found —", "raw": ""})
    else:
        matched_recipe_keys.add(key)
        ings = rec["ingredients"] or [("(no ingredients listed)", "")]
        for i,(ing,qty) in enumerate(ings):
            combined.append({**d,
                             "energy": d["energy"] if i==0 else None,
                             "carbs":  d["carbs"]  if i==0 else None,
                             "protein":d["protein"]if i==0 else None,
                             "fat":    d["fat"]    if i==0 else None,
                             "fiber":  d["fiber"]  if i==0 else None,
                             "ingredient": ing, "raw": qty})

# recipes never used in the weekly plan (orphans)
used_full = set()
for d in diet:
    used_full.add(norm(d["dish"]))
    for part in re.split(r'\s*&\s*', d["dish"]):
        used_full.add(norm(part))
orphan = sorted({v["name"] for k,v in recipes.items() if k not in used_full and v["name"] not in
                 [recipes[u]["name"] for u in used_full if u in recipes]})

# ---------- 4) Write combined workbook ----------
out = openpyxl.Workbook()
sh = out.active; sh.title = "Combined Menu"
HEAD = ["Day","Meal","Dish","Cooked Qty","Energy (kcal)","Carbs (g)","Protein (g)","Fat (g)","Fiber (g)","Ingredient","Raw Amt","Unit","Subtotal (g)","Subtotal (ml)"]
sh.append(HEAD)
hfill = PatternFill("solid", fgColor="1E3A5F"); hfont = Font(bold=True, color="FFFFFF", size=11)
thin = Side(style="thin", color="D9D9D9"); border = Border(thin,thin,thin,thin)
for j,_ in enumerate(HEAD,1):
    ccell = sh.cell(1,j); ccell.fill=hfill; ccell.font=hfont
    ccell.alignment=Alignment(vertical="center", wrap_text=True); ccell.border=border

def split_qty(name):
    """Isolate the trailing weight/quantity from a dish name.
    'Vegetable omlete-150g' -> ('Vegetable omlete','150g'); 'Banana- 2no' -> ('Banana','2no')."""
    n = str(name or "").strip()
    if "-" in n:
        left, right = n.rsplit("-", 1)
        if any(ch.isdigit() for ch in right):
            return left.strip(), right.strip()
    return n, ""

def split_raw(raw):
    """'200g' -> (200,'g'); '2.5g' -> (2.5,'g'); '100ml' -> (100,'ml')."""
    if not raw:
        return "", ""
    m = re.match(r'\s*([0-9]*\.?[0-9]+)\s*(.*)$', str(raw))
    if not m:
        return "", str(raw).strip()
    amt = float(m.group(1)); unit = m.group(2).strip()
    if amt == int(amt): amt = int(amt)
    return amt, unit

ZERO_BLANK = '0.###;-0.###;""'   # number format: hide zero, drop trailing zeros
from collections import OrderedDict
# Group flat rows into Day -> Meal -> Dish -> [ingredient rows], preserving order.
structure = OrderedDict()
for d in combined:
    structure.setdefault(d["day"], OrderedDict()).setdefault(d["meal"], OrderedDict()).setdefault(d["dish"], []).append(d)

meal_fill = PatternFill("solid", fgColor="E8F0FA")
day_fill  = PatternFill("solid", fgColor="C9DAF0")
dish_groups=[]   # (first,last) per dish for per-dish g/ml subtotals
band=False

def style_row(rr, fill=None, bold=False, color=None):
    for j in range(1,15):
        cc=sh.cell(rr,j); cc.border=border
        if fill: cc.fill=fill
        if j in (5,6,7,8,9,11,13,14) or j==4: cc.alignment=Alignment(horizontal="right")
        if j==12: cc.alignment=Alignment(horizontal="center")
        if bold: cc.font=Font(bold=True, color=color or "1E3A5F")

for day, meals in structure.items():
    day_started=False
    day_meal_sub_rows=[]        # meal-subtotal row numbers for this day
    day_dish_first=None; day_dish_last=None
    for meal, dishes in meals.items():
        meal_started=False
        meal_first=sh.max_row+1
        for dish, rows in dishes.items():
            band = not band
            dish_first=sh.max_row+1
            for i,d in enumerate(rows):
                amt_v, unit_v = split_raw(d["raw"])
                dish_name, cooked_qty = split_qty(d["dish"])
                day_v  = day  if not day_started  else ""
                meal_v = meal if not meal_started else ""
                dish_v = dish_name  if i==0 else ""
                qty_v  = cooked_qty if i==0 else ""
                day_started=True; meal_started=True
                nut = (d["energy"], d["carbs"], d["protein"], d["fat"], d["fiber"]) if i==0 else (None,)*5
                sh.append([day_v, meal_v, dish_v, qty_v, *nut, d["ingredient"], amt_v, unit_v, None, None])
                rr=sh.max_row
                style_row(rr, fill=PatternFill("solid", fgColor="FFFFFF" if band else "F5F8FB"))
                if i==0 and dish_v: sh.cell(rr,3).font=Font(bold=True)
                if d["ingredient"].startswith("— no recipe"):
                    for j in range(1,15): sh.cell(rr,j).font=Font(color="C0392B", italic=True)
            dish_groups.append((dish_first, sh.max_row))
        meal_last=sh.max_row
        if day_dish_first is None: day_dish_first=meal_first
        day_dish_last=meal_last
        # Meal subtotal row (skip when meal label is blank, e.g. Extras)
        if str(meal).strip():
            sh.append(["","",f"{meal} — subtotal","",
                       f"=SUM(E{meal_first}:E{meal_last})", f"=SUM(F{meal_first}:F{meal_last})",
                       f"=SUM(G{meal_first}:G{meal_last})", f"=SUM(H{meal_first}:H{meal_last})",
                       f"=SUM(I{meal_first}:I{meal_last})", "","","",
                       f"=SUM(M{meal_first}:M{meal_last})", f"=SUM(N{meal_first}:N{meal_last})"])
            sr=sh.max_row
            style_row(sr, fill=meal_fill, bold=True)
            for c in (13,14): sh.cell(sr,c).number_format=ZERO_BLANK
            day_meal_sub_rows.append(sr)
    # Day TOTAL row — sum the meal-subtotal rows (avoids double counting);
    # for Extras (no meal subtotals) sum the day's dish range directly.
    if day_meal_sub_rows:
        mk = lambda col: "=SUM(" + ",".join(f"{col}{r}" for r in day_meal_sub_rows) + ")"
    else:
        mk = lambda col: f"=SUM({col}{day_dish_first}:{col}{day_dish_last})"
    sh.append(["","",f"{day} — DAY TOTAL","", mk("E"),mk("F"),mk("G"),mk("H"),mk("I"),"","","", mk("M"),mk("N")])
    dr=sh.max_row
    style_row(dr, fill=day_fill, bold=True, color="0B2545")
    sh.cell(dr,3).font=Font(bold=True, color="0B2545", size=11)
    for c in (13,14): sh.cell(dr,c).number_format=ZERO_BLANK

# Live per-dish subtotals on each dish's first row — grams (M) & ml (N)
# summed SEPARATELY via SUMIFS over the visible Raw Amt + Unit columns.
for first, last in dish_groups:
    for col, unit in ((13,"g"), (14,"ml")):
        c = sh.cell(first, col)
        c.value = f'=SUMIFS(K{first}:K{last},L{first}:L{last},"{unit}")'
        c.font = Font(bold=True, color="1E3A5F"); c.alignment = Alignment(horizontal="right")
        c.number_format = ZERO_BLANK

widths=[16,14,30,11,12,10,11,9,9,22,9,7,12,12]
for i,w in enumerate(widths,1):
    sh.column_dimensions[openpyxl.utils.get_column_letter(i)].width=w
sh.freeze_panes="A2"

# Mismatches sheet
ms = out.create_sheet("Mismatches & Notes")
ms.append(["Kitchen menu — combine report"]); ms["A1"].font=Font(bold=True,size=13)
ms.append([])
ms.append([f"Weekly-diet dishes with NO matching recipe ({len(set(diet_no_recipe))}):"])
ms["A3"].font=Font(bold=True,color="C0392B")
for n in sorted(set(diet_no_recipe)): ms.append(["", n])
ms.append([])
r0=ms.max_row+1
ms.append([f"Recipes never used in the weekly plan ({len(orphan)}):"])
ms.cell(r0, 1).font=Font(bold=True, color="B9770E")
for n in orphan: ms.append(["", n])
ms.column_dimensions["A"].width=6; ms.column_dimensions["B"].width=60

out.save('/app/frontend/public/combined_menu.xlsx')

# ---------- report to console ----------
print("diet dishes (rows):", len(diet))
print("recipe dishes:", len({v['name'] for v in recipes.values()}))
print("combined rows:", len(combined))
print("\nDIET DISHES WITH NO RECIPE (%d):" % len(set(diet_no_recipe)))
for n in sorted(set(diet_no_recipe)): print("  -", n)
print("\nORPHAN RECIPES not in weekly plan (%d):" % len(orphan))
for n in orphan: print("  -", n)
