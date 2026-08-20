"""One-shot builder for the combined Menu Master workbook.

Takes:
  • Calculations of menu for the app.xlsx  (menu → kcal / macros)
  • Recipes of the menu (major ingredients).xlsx  (menu → ingredients)

and writes a 5-sheet consolidated workbook to
`/app/backend/data/menu_master_combined.xlsx`, which is served by
`GET /api/meals/menu-master-download`.

Rerun whenever the source spreadsheets are updated:
    python -m scripts.build_menu_master \\
        /path/to/calculations.xlsx /path/to/recipes.xlsx
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

import pandas as pd

DAYS = ["Monday", "Tuesday", "Wednesday", "Thursday",
        "Friday", "Saturday", "Sunday", "Elite Extras"]
MEAL_ORDER = ["Pre-Training", "Pre -Breakfast", "Breakfast", "Mid morning",
              "Lunch", "Post-Training", "Snack", "Dinner", "Elite Extra", "DAY TOTAL"]

DEFAULT_CALC = Path(__file__).resolve().parents[1] / "data" / "menu_calcs.xlsx"
DEFAULT_RECIPES = Path(__file__).resolve().parents[1] / "data" / "recipes.xlsx"
OUT_PATH = Path(__file__).resolve().parents[1] / "data" / "menu_master_combined.xlsx"


def _parse_calc_block(df, col_start, day, has_meal_col):
    rows = []; cur_meal = None
    for r in range(3, len(df)):
        row = df.iloc[r]
        if has_meal_col:
            meal_cell = str(row.iloc[col_start]).strip() if pd.notna(row.iloc[col_start]) else ""
            if meal_cell and meal_cell.lower() != "nan":
                cur_meal = meal_cell
            menu_cell = row.iloc[col_start + 1]
            vals = [row.iloc[col_start + i] for i in range(2, 7)]
        else:
            cur_meal = "Elite Extra"
            menu_cell = row.iloc[col_start]
            vals = [row.iloc[col_start + i] for i in range(1, 5)] + [None]
        if pd.isna(menu_cell) or not str(menu_cell).strip():
            continue
        rows.append({
            "day": day, "meal": cur_meal,
            "menu_item": str(menu_cell).strip(),
            "kcal":      float(vals[0]) if pd.notna(vals[0]) else None,
            "carbs_g":   float(vals[1]) if pd.notna(vals[1]) else None,
            "protein_g": float(vals[2]) if pd.notna(vals[2]) else None,
            "fat_g":     float(vals[3]) if pd.notna(vals[3]) else None,
            "fibre_g":   float(vals[4]) if pd.notna(vals[4]) else None,
        })
    return rows


def _parse_recipes_block(df, col_start, force_day=None):
    rows = []; cur_day = force_day; cur_meal = None; cur_menu = None
    for r in range(1, len(df)):
        row = df.iloc[r]
        cell = row.iloc[col_start]
        if not force_day and pd.notna(cell) and "Day-" in str(cell):
            cur_day = str(cell).replace("Day-", "").strip()
            continue
        if pd.notna(cell) and str(cell).strip() == "Meal":
            continue
        meal_cell = str(cell).strip() if pd.notna(cell) else ""
        if meal_cell and meal_cell.lower() != "nan":
            cur_meal = meal_cell
        menu_cell = row.iloc[col_start + 1] if col_start + 1 < len(row) else None
        if pd.notna(menu_cell) and str(menu_cell).strip():
            cur_menu = str(menu_cell).strip()
        ing_cell = row.iloc[col_start + 2] if col_start + 2 < len(row) else None
        qty_cell = row.iloc[col_start + 3] if col_start + 3 < len(row) else None
        if pd.isna(ing_cell) or not str(ing_cell).strip() or str(ing_cell).strip().lower() == "nan":
            continue
        rows.append({
            "day": cur_day, "meal": cur_meal, "menu_item": cur_menu,
            "ingredient": str(ing_cell).strip(),
            "raw_qty": str(qty_cell).strip() if pd.notna(qty_cell) else "",
        })
    return rows


def _split_menu(m):
    if not isinstance(m, str):
        return m, None
    parts = re.split(r"-\s*", m.strip(), maxsplit=1)
    return (parts[0].strip(), parts[1].strip()) if len(parts) == 2 else (m.strip(), None)


def _to_grams(qty):
    if not qty:
        return None
    s = str(qty).lower().strip().replace("approx", "").strip()
    m = re.match(r"^\s*([\d\.]+)\s*(kg|g|ml|l|no|nos|pc|pcs|slice|slices|tsp|tbsp|teaspoon|tablespoon|cup|small|medium|large)?", s)
    if not m:
        return None
    n = float(m.group(1))
    u = (m.group(2) or "g").lower()
    if u == "kg":  return n * 1000
    if u in ("ml", "l"):
        return n * (1000 if u == "l" else 1)  # 1 ml ≈ 1 g for cooking liquids
    return n


def build(calc_path: Path, recipes_path: Path, out_path: Path = OUT_PATH) -> None:
    raw1 = pd.read_excel(calc_path, sheet_name="Sheet1", header=None)
    raw2 = pd.read_excel(recipes_path, sheet_name="Sheet1", header=None)

    calc_rows = []
    for d, cs in zip(DAYS[:7], [0, 8, 16, 24, 32, 40, 48]):
        calc_rows.extend(_parse_calc_block(raw1, cs, d, True))
    calc_rows.extend(_parse_calc_block(raw1, 56, "Elite Extras", False))
    calc_df = pd.DataFrame(calc_rows)

    recipe_rows = _parse_recipes_block(raw2, 0)
    recipe_rows.extend(_parse_recipes_block(raw2, 5, force_day="Elite Extras"))
    recipes_df = pd.DataFrame(recipe_rows)

    for df in (calc_df, recipes_df):
        df[["menu_name", "cooked_qty"]] = df["menu_item"].apply(lambda x: pd.Series(_split_menu(x)))
    recipes_df["raw_grams"] = recipes_df["raw_qty"].apply(_to_grams)

    merged = recipes_df.merge(
        calc_df[["day", "menu_item", "kcal", "carbs_g", "protein_g", "fat_g", "fibre_g"]],
        on=["day", "menu_item"], how="left",
    )
    # Fallback: match on menu_name only when exact join failed
    missing = merged["kcal"].isna()
    if missing.any():
        by_name = calc_df.groupby(["day", "menu_name"])[
            ["kcal", "carbs_g", "protein_g", "fat_g", "fibre_g"]].first().reset_index()
        fill = merged[missing].drop(columns=["kcal", "carbs_g", "protein_g", "fat_g", "fibre_g"]) \
                              .merge(by_name, on=["day", "menu_name"], how="left")
        for col in ["kcal", "carbs_g", "protein_g", "fat_g", "fibre_g"]:
            merged.loc[missing, col] = fill[col].values

    master = merged[["day", "meal", "menu_name", "cooked_qty", "ingredient", "raw_qty", "raw_grams",
                     "kcal", "carbs_g", "protein_g", "fat_g", "fibre_g"]].copy()
    master.columns = ["Day", "Meal", "Menu Item", "Cooked Qty", "Ingredient", "Raw Qty", "Raw grams",
                      "Menu kcal", "Menu Carbs (g)", "Menu Protein (g)", "Menu Fat (g)", "Menu Fibre (g)"]

    menu_summary_rows = []
    for (day, meal, menu_name), group in master.groupby(["Day", "Meal", "Menu Item"]):
        menu_summary_rows.append({
            "Day": day, "Meal": meal, "Menu Item": menu_name,
            "Cooked Qty": group["Cooked Qty"].iloc[0],
            "Ingredients": " + ".join(f"{i} ({r})" for i, r in zip(group["Ingredient"], group["Raw Qty"])),
            "Total Raw (g)": group["Raw grams"].fillna(0).sum() or None,
            "kcal": group["Menu kcal"].iloc[0],
            "Carbs (g)": group["Menu Carbs (g)"].iloc[0],
            "Protein (g)": group["Menu Protein (g)"].iloc[0],
            "Fat (g)": group["Menu Fat (g)"].iloc[0],
            "Fibre (g)": group["Menu Fibre (g)"].iloc[0],
        })
    menu_summary = pd.DataFrame(menu_summary_rows)

    uniq = menu_summary.drop_duplicates(["Day", "Meal", "Menu Item"])
    day_totals = uniq.groupby(["Day", "Meal"])[
        ["kcal", "Carbs (g)", "Protein (g)", "Fat (g)", "Fibre (g)"]].sum().reset_index()
    grand = uniq.groupby(["Day"])[
        ["kcal", "Carbs (g)", "Protein (g)", "Fat (g)", "Fibre (g)"]].sum().reset_index()
    grand["Meal"] = "DAY TOTAL"
    day_totals = pd.concat([day_totals, grand[day_totals.columns]], ignore_index=True)
    day_totals["_dix"] = day_totals["Day"].apply(lambda x: DAYS.index(x) if x in DAYS else 99)
    day_totals["_mix"] = day_totals["Meal"].apply(lambda m: MEAL_ORDER.index(m) if m in MEAL_ORDER else 50)
    day_totals = day_totals.sort_values(["_dix", "_mix"]).drop(columns=["_dix", "_mix"]).reset_index(drop=True)

    with pd.ExcelWriter(out_path, engine="openpyxl") as xw:
        master.to_excel(xw, sheet_name="Master (ingredient level)", index=False)
        menu_summary.to_excel(xw, sheet_name="Menu Summary", index=False)
        day_totals.to_excel(xw, sheet_name="Day Totals", index=False)
        calc_df.rename(columns={
            "day": "Day", "meal": "Meal", "menu_item": "Menu Item",
            "menu_name": "Menu Name", "cooked_qty": "Cooked Qty",
            "carbs_g": "Carbs (g)", "protein_g": "Protein (g)",
            "fat_g": "Fat (g)", "fibre_g": "Fibre (g)",
        }).to_excel(xw, sheet_name="Menu-Nutrition (raw)", index=False)
        recipes_df.rename(columns={
            "day": "Day", "meal": "Meal", "menu_item": "Menu Item",
            "menu_name": "Menu Name", "cooked_qty": "Cooked Qty",
            "ingredient": "Ingredient", "raw_qty": "Raw Qty",
            "raw_grams": "Raw grams",
        }).to_excel(xw, sheet_name="Recipes (raw)", index=False)

    # Cosmetic touch-up: bold slate header + freeze first row + auto-width.
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = load_workbook(out_path)
    for sn in wb.sheetnames:
        ws = wb[sn]
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1E293B")
            cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.freeze_panes = "A2"
        for col in ws.columns:
            try:
                w = max((len(str(c.value)) for c in col if c.value is not None), default=10) + 2
                ws.column_dimensions[col[0].column_letter].width = min(w, 45)
            except Exception:
                pass
    wb.save(out_path)

    print(f"✓ Master workbook written to {out_path}")


def main() -> None:
    argv = sys.argv[1:]
    calc = Path(argv[0]) if len(argv) > 0 else DEFAULT_CALC
    recipes = Path(argv[1]) if len(argv) > 1 else DEFAULT_RECIPES
    if not calc.exists() or not recipes.exists():
        print(f"✗ source file(s) missing:\n  calc={calc}\n  recipes={recipes}", file=sys.stderr)
        sys.exit(1)
    build(calc, recipes)


if __name__ == "__main__":
    main()
